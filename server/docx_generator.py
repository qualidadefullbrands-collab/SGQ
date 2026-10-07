from __future__ import annotations
from io import BytesIO
from typing import Any
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from PIL import Image
import re

def fmt_date(value: str | None) -> str:
    if not value: return ""
    p=str(value)[:10].split("-")
    return f"{p[2]}/{p[1]}/{p[0]}" if len(p)==3 else str(value)

def pct(n:int,d:int):
    return round((n/d)*100,1) if d else None

def fmt_pct(v):
    return "" if v is None else f"{v:.1f}%".replace(".",",")

def classify(v):
    if v is None: return ""
    if v < 60: return "Insatisfatório"
    if v < 75: return "Regular"
    if v < 80: return "Bom"
    return "Excelente"

def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"): p._p.remove(child)

def set_font(run,size=8,bold=False):
    run.font.name="Calibri"
    if run._element.rPr is None: run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:eastAsia"),"Calibri")
    run.font.size=Pt(size)
    run.bold=bold

def set_cell(cell,text:Any,size=8,bold=False,center=False):
    text="" if text is None else str(text)
    p=cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    clear_paragraph(p)
    r=p.add_run(text); set_font(r,size,bold)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p2 in list(cell.paragraphs[1:]):
        p2._element.getparent().remove(p2._element)

def add_row(table,values,size=7.5):
    row=table.add_row()
    for i,v in enumerate(values):
        if i < len(row.cells): set_cell(row.cells[i],v,size)
    return row

def photo_bytes(source):
    return source if isinstance(source,(bytes,bytearray)) else None

def add_photo(cell,data:bytes,caption:str,width=1.75):
    im=Image.open(BytesIO(data)); im.thumbnail((1100,800))
    out=BytesIO(); im.convert("RGB").save(out,format="JPEG",quality=84); out.seek(0)
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(out,width=Inches(width))
    cap=cell.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=cap.add_run(caption); set_font(r,7)

def fill_photos(table,photos,max_photos):
    for idx,item in enumerate(photos[:max_photos]):
        data=photo_bytes(item.get("bytes"))
        if not data: continue
        add_photo(table.cell(idx//3,idx%3),data,f"Foto {idx+1}: {str(item.get('caption') or '').strip()}")

def c_count(a): return sum(1 for x in a if x.get("resultado")=="C")
def nc_count(a): return sum(1 for x in a if x.get("resultado")=="NC")
def na_count(a): return sum(1 for x in a if x.get("resultado")=="NA")

def road(local):
    s=str(local or "").strip()
    m=re.search(r"\bRua\s*([0-9]{1,3})\b",s,re.I)
    return f"Rua {m.group(1).zfill(2)}" if m else s

def generate_rq016(doc,p):
    ex=p["execution"]; a=p.get("answers") or []; f=p.get("findings") or []
    t0,t1,t2,t3,t4,t5,t6=doc.tables
    set_cell(t0.cell(0,1),fmt_date(ex.get("data_avaliacao"))); set_cell(t0.cell(0,3),ex.get("responsavel_nome",""))
    c,nc,na=c_count(a),nc_count(a),na_count(a); note=pct(c,c+nc)
    for i,v in enumerate([c,nc,na,fmt_pct(note),classify(note)]): set_cell(t1.cell(1,i),v,bold=True,center=True)
    bykey={str(x.get("item_key")):x for x in a}
    for rowx in t2.rows[1:]:
        k=rowx.cells[0].text.strip()
        if k.isdigit() and k in bykey: set_cell(rowx.cells[2],bykey[k].get("resultado",""),bold=True,center=True)
    for i,x in enumerate(f,1): add_row(t3,[i,x.get("descricao_tecnica") or x.get("descricao_original",""),x.get("acao_corretiva",""),x.get("responsavel",""),fmt_date(x.get("prazo")),x.get("status","Pendente")])
    for i,x in enumerate(p.get("previous") or [],1): add_row(t4,[i,x.get("descricao",""),fmt_date(x.get("aberta_em")),x.get("situacao",""),x.get("status","")])
    fill_photos(t5,p.get("photos") or [],18)
    set_cell(t6.cell(1,0),ex.get("responsavel_nome","")); set_cell(t6.cell(1,1),ex.get("validado_por","")); set_cell(t6.cell(1,2),fmt_date(ex.get("data_validacao")))

def generate_rq016b(doc,p):
    ex=p["execution"]; a=p.get("answers") or []; f=p.get("findings") or []
    t0,t1,_,_,_,t5,t6,t7=doc.tables
    set_cell(t0.cell(0,1),fmt_date(ex.get("data_avaliacao"))); set_cell(t0.cell(0,3),ex.get("mes_referencia",""))
    set_cell(t0.cell(1,1),ex.get("responsavel_nome",""))
    locs=sorted({str(x.get("local_ref") or "").strip() for x in a if x.get("local_ref")}); set_cell(t0.cell(1,3),", ".join(locs),7.5)
    def area(k):
        k=str(k or "")
        return "Ruas" if k.startswith("R") else "Checkouts" if k.startswith("C") else "Docas" if k.startswith("D") else ""
    areas={x:[] for x in ["Ruas","Checkouts","Docas"]}
    for x in a:
        ar=area(x.get("item_key"))
        if ar: areas[ar].append(x)
    def score(items):
        c=sum(1 for x in items if x.get("resultado")=="C"); nc=sum(1 for x in items if x.get("resultado")=="NC")
        return pct(c,c+nc)
    total=score(a)
    vals=[fmt_pct(total),fmt_pct(score(areas["Ruas"])),fmt_pct(score(areas["Checkouts"])),fmt_pct(score(areas["Docas"])),classify(total)]
    for i,v in enumerate(vals): set_cell(t1.cell(1,i),v,bold=True,center=True)
    byloc={}
    for x in a: byloc.setdefault(x.get("local_ref",""),[]).append(x)
    for loc,items in sorted(byloc.items()):
        ncs=[str(x.get("item_key")) for x in items if x.get("resultado")=="NC"]
        obs="; ".join([str(x.get("observacao") or "") for x in items if x.get("resultado")=="NC" and x.get("observacao")])[:260]
        add_row(t5,[area(items[0].get("item_key")) if items else "",loc,fmt_pct(score(items)),", ".join(ncs),obs])
    for x in f: add_row(t6,[x.get("local_ref") or x.get("endereco",""),x.get("descricao_tecnica") or x.get("descricao_original",""),x.get("acao_corretiva",""),x.get("responsavel",""),fmt_date(x.get("prazo")),x.get("status","Pendente")])
    set_cell(t7.cell(1,0),ex.get("responsavel_nome","")); set_cell(t7.cell(1,1),ex.get("validado_por","")); set_cell(t7.cell(1,2),fmt_date(ex.get("data_validacao")))

def generate_rq015(doc,p):
    ex=p["execution"]; a=p.get("answers") or []; f=p.get("findings") or []
    t0,t1,_,_,t4,t5,t6,t7,t8,t9=doc.tables
    set_cell(t0.cell(0,1),fmt_date(ex.get("data_avaliacao"))); set_cell(t0.cell(0,3),ex.get("mes_referencia",""))
    set_cell(t0.cell(1,1),ex.get("responsavel_nome","")); set_cell(t0.cell(1,3),ex.get("documento_vinculado",""))
    roads=sorted({road(x.get("local_ref")) for x in a if x.get("local_ref")} | {road(x.get("local_ref")) for x in f if x.get("local_ref")})
    risks={"verde":0,"amarelo":0,"vermelho":0}
    for x in f:
        r=str(x.get("risco") or "").lower()
        if r in risks: risks[r]+=1
    for i,v in enumerate([len(roads),len(f),risks["verde"],risks["amarelo"],risks["vermelho"],risks["vermelho"]]): set_cell(t1.cell(1,i),v,bold=True,center=True)
    byroad={r:[] for r in roads}
    for x in a:
        r=road(x.get("local_ref"))
        if r: byroad.setdefault(r,[]).append(x)
    for r in roads:
        aa=byroad.get(r,[]); ncs=[x for x in aa if x.get("resultado")=="NC"]
        result="Não conforme" if ncs else ("Conforme" if aa else "")
        count=sum(1 for x in f if road(x.get("local_ref"))==r)
        obs="; ".join([str(x.get("observacao") or "") for x in ncs if x.get("observacao")])[:300]
        add_row(t4,[r,result,count,obs])
    for i,x in enumerate(f,1):
        add_row(t5,[i,x.get("endereco") or x.get("local_ref",""),", ".join(x.get("criterios") or []),x.get("descricao_tecnica") or x.get("descricao_original",""),str(x.get("risco") or "").capitalize(),x.get("acao_imediata",""),x.get("status","Pendente")])
        if x.get("acao_corretiva") or x.get("responsavel") or x.get("prazo"):
            meta=x.get("metadata") or {}
            add_row(t6,[i,x.get("acao_corretiva",""),x.get("responsavel",""),fmt_date(x.get("prazo")),meta.get("os_rq030",""),meta.get("conclusao","")])
    for x in p.get("previous") or []: add_row(t7,[x.get("endereco",""),x.get("descricao",""),fmt_date(x.get("aberta_em")),x.get("situacao",""),x.get("status","")])
    fill_photos(t8,p.get("photos") or [],24)
    set_cell(t9.cell(1,0),ex.get("responsavel_nome","")); set_cell(t9.cell(1,1),ex.get("validado_por","")); set_cell(t9.cell(1,2),fmt_date(ex.get("data_validacao")))

def generate_rq014(doc,p):
    ex=p["execution"]; f=p.get("findings") or []; objects=p.get("objects") or []
    t0,t1,_,t3,_,t5,t6,t7,t8,t9=doc.tables
    set_cell(t0.cell(0,1),fmt_date(ex.get("data_avaliacao"))); set_cell(t0.cell(0,3),ex.get("mes_referencia",""))
    set_cell(t0.cell(1,1),ex.get("responsavel_nome","")); set_cell(t0.cell(1,3),ex.get("documento_vinculado",""))
    lib=sum(1 for x in objects if str(x.get("status") or "").lower()=="liberado"); pend=sum(1 for x in objects if str(x.get("status") or "").lower()=="pendente"); bloq=sum(1 for x in objects if str(x.get("status") or "").lower()=="bloqueado")
    for i,v in enumerate([len(objects),lib,pend,bloq,fmt_pct(pct(lib,len(objects)))]): set_cell(t1.cell(1,i),v,bold=True,center=True)
    ext=[x for x in objects if x.get("objeto_tipo")=="extintor"]; hid=[x for x in objects if x.get("objeto_tipo")=="hidrante"]
    for i,x in enumerate(ext,1):
        d=x.get("dados") or {}; add_row(t3,[i,x.get("identificacao",""),x.get("local_ref",""),d.get("tipo_capacidade",""),d.get("recarga_ate",""),d.get("hidro_ate",""),", ".join(x.get("criterios_nc") or []),str(x.get("status") or "").capitalize(),d.get("os_rq030","")])
    for i,x in enumerate(hid,1):
        d=x.get("dados") or {}; add_row(t5,[i,x.get("identificacao",""),x.get("local_ref",""),d.get("mangueira",""),", ".join(x.get("criterios_nc") or []),str(x.get("status") or "").capitalize(),d.get("os_rq030","")])
    for x in f:
        meta=x.get("metadata") or {}; add_row(t6,[x.get("endereco") or x.get("local_ref",""),x.get("descricao_tecnica") or x.get("descricao_original",""),x.get("acao_corretiva",""),x.get("responsavel",""),fmt_date(x.get("prazo")),meta.get("os_rq030",""),x.get("status","Pendente")])
    for x in p.get("previous") or []: add_row(t7,[x.get("item",""),x.get("descricao",""),fmt_date(x.get("aberta_em")),x.get("situacao",""),x.get("status","")])
    fill_photos(t8,p.get("photos") or [],18)
    set_cell(t9.cell(1,0),ex.get("responsavel_nome","")); set_cell(t9.cell(1,1),ex.get("validado_por","")); set_cell(t9.cell(1,2),fmt_date(ex.get("data_validacao")))

def generate(template_bytes:bytes,payload:dict)->bytes:
    doc=Document(BytesIO(template_bytes)); rq=payload["execution"]["rq_code"]
    {"RQ014":generate_rq014,"RQ015":generate_rq015,"RQ016":generate_rq016,"RQ016B":generate_rq016b}[rq](doc,payload)
    out=BytesIO(); doc.save(out); return out.getvalue()

MONTHS=["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO","AGOSTO","SETEMBRO","OUTUBRO","NOVEMBRO","DEZEMBRO"]
def file_name(rq:str,date_raw:str)->str:
    from datetime import date
    d=date.fromisoformat(str(date_raw)[:10]); month=MONTHS[d.month-1]
    if rq=="RQ014": return f"RQ 014 - {month}.{d.year}.docx"
    if rq=="RQ015": return f"RQ 015 - {month}.docx"
    if rq=="RQ016": return f"RQ 016 - Avaliação 5S {d.day:02d}.{d.month:02d}.{d.year}.docx"
    return f"Avaliação Pré Avaliação 5S - {month.capitalize()}.{d.year}.docx"


def _norm(value):
    return re.sub(r"\s+"," ",str(value or "")).strip().lower()

def _mark_choice(text, choice):
    raw=str(text or "")
    labels={
        "I":"Nível I","II":"Nível II","III":"Nível III",
        "S1":"Especial S1","S2":"Especial S2","S3":"Especial S3","S4":"Especial S4",
        "aprovado":"APROVADO","reprovado":"REPROVADO",
    }
    target=labels.get(str(choice),str(choice))
    def repl(m):
        label=m.group(2)
        return f"( X ) {label}" if _norm(label)==_norm(target) else f"(   ) {label}"
    return re.sub(r"\(\s*[xX]?\s*\)\s*([^()]+?)(?=(?:\s*\(\s*[xX]?\s*\))|$)",lambda m: repl(type("M",(),{"group":lambda self,n: m.group(0) if n==0 else (None if n==1 else m.group(1))})()),raw)

def _fill_choice_text(raw, selected):
    raw=str(raw or "")
    candidates=["Nível I","Nível II","Nível III","Especial S1","Especial S2","Especial S3","Especial S4","APROVADO","REPROVADO"]
    out=raw
    for label in candidates:
        if label.lower() not in out.lower():
            continue
        pattern=re.compile(r"\(\s*[xX]?\s*\)\s*"+re.escape(label),re.I)
        out=pattern.sub(("( X ) " if _norm(label)==_norm(selected) else "(   ) ")+label,out)
    return out

def _first(payload,key,default=""):
    v=payload.get(key,default)
    return "" if v is None else str(v)

def generate_inspection(template_bytes:bytes,payload:dict)->bytes:
    doc=Document(BytesIO(template_bytes))
    inspection=payload.get("inspection") or {}
    process=payload.get("process") or {}
    items=payload.get("items") or []
    checks=payload.get("checklist") or []
    check_results=payload.get("check_results") or []
    dim_params=payload.get("dim_params") or []
    dim_configs=payload.get("dim_configs") or []
    dim_results=payload.get("dim_results") or []
    photos=payload.get("photos") or []
    inspector=payload.get("inspector_name") or ""

    first_item=items[0] if items else {}
    product=first_item.get("produto") or {}
    identification={
        "produto / descrição": product.get("nome") or "",
        "código do cliente": product.get("sku") or "",
        "cliente": process.get("cliente") or "",
        "processo fst": process.get("codigo") or "",
        "nota fiscal": process.get("nota_fiscal") or "",
        "lote": first_item.get("lote") or "",
        "material": first_item.get("material") or "",
        "capacidade": first_item.get("capacidade") or "",
        "origem": process.get("origem") or "",
        "transporte": process.get("transporte") or "",
        "quantidade recebida": first_item.get("quantidade") or "",
        "quantidade por caixa": first_item.get("quantidade_por_caixa") or "",
        "nº de caixas recebidas": first_item.get("caixas_recebidas") or "",
        "n° de caixas recebidas": first_item.get("caixas_recebidas") or "",
        "chegada no cd": fmt_date(process.get("chegada_cd")),
        "data da inspeção": fmt_date(inspection.get("data_inspecao")),
    }

    for table in doc.tables:
        for row in table.rows:
            cells=row.cells
            if len(cells)>=2:
                for idx in range(0,len(cells)-1):
                    label=_norm(cells[idx].text)
                    if label in identification:
                        set_cell(cells[idx+1],identification[label],8)
            joined=" ".join(c.text for c in cells)
            if "Nível de inspeção" in joined or "Nível I" in joined:
                level=str(inspection.get("nivel_inspecao") or "")
                for cell in cells:
                    if "Nível" in cell.text or "Especial" in cell.text:
                        set_cell(cell,_fill_choice_text(cell.text,{"I":"Nível I","II":"Nível II","III":"Nível III","S1":"Especial S1","S2":"Especial S2","S3":"Especial S3","S4":"Especial S4"}.get(level,level)),8)
            if len(cells)>=2 and _norm(cells[0].text)=="quantidade amostrada":
                set_cell(cells[1],inspection.get("tamanho_amostra") or inspection.get("total_inspecionado") or "",8,bold=True,center=True)
            if len(cells)>=4 and _norm(cells[2].text)=="caixas avaliadas":
                set_cell(cells[3],inspection.get("caixas_avaliar") or "",8,bold=True,center=True)

    by_check={str(x.get("checklist_id")):x for x in check_results}
    by_order={str(x.get("ordem")):x for x in checks}
    for table in doc.tables:
        if not table.rows: continue
        headers=[_norm(c.text) for c in table.rows[0].cells]
        if "análise" in headers and "c" in headers and "nc" in headers and "na" in headers:
            for row in table.rows[1:]:
                order=row.cells[0].text.strip().rstrip(".")
                check=by_order.get(order)
                if not check: continue
                result=by_check.get(str(check.get("id")),{}).get("resultado")
                if len(row.cells)>=6:
                    set_cell(row.cells[3],"X" if result=="conforme" else "",9,bold=True,center=True)
                    set_cell(row.cells[4],"X" if result=="nao_conforme" else "",9,bold=True,center=True)
                    set_cell(row.cells[5],"X" if result=="nao_aplicavel" else "",9,bold=True,center=True)

    config_key={(str(x.get("processo_item_id")),str(x.get("parametro_id"))):x for x in dim_configs}
    result_key={}
    for x in dim_results:
        result_key.setdefault((str(x.get("processo_item_id")),str(x.get("parametro_id"))),[]).append(x)
    for vals in result_key.values():
        vals.sort(key=lambda x:int(x.get("sequencia_amostra") or 0))

    first_item_id=str(first_item.get("id") or "")
    params_by_name={_norm(x.get("nome")):x for x in dim_params}
    for table in doc.tables:
        for ridx,row in enumerate(table.rows):
            if len(row.cells)<4: continue
            pname=_norm(row.cells[0].text)
            param=params_by_name.get(pname)
            if not param: continue
            pid=str(param.get("id"))
            cfg=config_key.get((first_item_id,pid),{})
            vals=result_key.get((first_item_id,pid),[])
            if len(row.cells)>1:
                base=row.cells[1].text.split(":")[0].strip()
                set_cell(row.cells[1],f"{base}: {cfg.get('equipamento') or ''}",7.5)
            if len(row.cells)>2:
                base=row.cells[2].text.split(":")[0].strip()
                set_cell(row.cells[2],f"{base}: {cfg.get('codigo_equipamento') or ''}",7.5)
            if len(row.cells)>3:
                base=row.cells[3].text.split(":")[0].strip()
                spec=cfg.get("especificacao_desvio") or ""
                set_cell(row.cells[3],f"{base}: {spec}",7.5)
            for rr in table.rows[ridx+1:min(ridx+3,len(table.rows))]:
                for cell in rr.cells:
                    m=re.match(r"\s*(\d+)\.",cell.text)
                    if not m: continue
                    seq=int(m.group(1))
                    hit=next((x for x in vals if int(x.get("sequencia_amostra") or 0)==seq),None)
                    if hit:
                        value=hit.get("valor")
                        unit=hit.get("unidade") or cfg.get("unidade") or param.get("unidade") or ""
                        set_cell(cell,f"{seq}. {value} {unit}".strip(),7.5)

    any_dim_nc=any(x.get("conforme") is False for x in dim_results)
    for table in doc.tables:
        for row in table.rows:
            if row.cells and _norm(row.cells[0].text)=="resultado dimensional":
                for cell in row.cells[1:]:
                    set_cell(cell,_fill_choice_text(cell.text,"REPROVADO" if any_dim_nc else "APROVADO"),8,bold=True)

    final_result=str(inspection.get("resultado") or "").lower()
    for table in doc.tables:
        for row in table.rows:
            if not row.cells: continue
            label=_norm(row.cells[0].text)
            if label=="resultado final":
                for cell in row.cells[1:]:
                    set_cell(cell,_fill_choice_text(cell.text,"REPROVADO" if final_result=="reprovado" else "APROVADO"),8,bold=True)
            elif label=="inspecionado por" and len(row.cells)>=2:
                set_cell(row.cells[1],inspector,8)
                if len(row.cells)>=4: set_cell(row.cells[3],fmt_date(inspection.get("data_inspecao")),8)
            elif label.startswith("responsável técnico") and len(row.cells)>=2:
                set_cell(row.cells[1],"Vanessa T. Casarin – Responsável Técnica – CRF 22864",8)
            elif label.startswith("observação / conclusão") and len(row.cells)>=2:
                set_cell(row.cells[1],inspection.get("observacoes") or "",8)

    photo_table=None
    for idx,table in enumerate(doc.tables):
        if any("FOTO" in (c.text or "").upper() for row in table.rows for c in row.cells):
            if len(table.rows)>=2 and len(table.columns)>=2:
                photo_table=table
    if photo_table and photos:
        cells=[cell for row in photo_table.rows for cell in row.cells]
        for idx,item in enumerate(photos[:len(cells)]):
            data=photo_bytes(item.get("bytes"))
            if data:
                add_photo(cells[idx],data,item.get("caption") or "",width=1.75)

    out=BytesIO(); doc.save(out); return out.getvalue()

def inspection_file_name(process_code:str|None,inspection_number:str|None)->str:
    process=str(process_code or "FST").replace("/","-")
    number=str(inspection_number or "INSPECAO").replace("/","-")
    return f"{process} - {number} - Laudo de Inspeção.docx"
